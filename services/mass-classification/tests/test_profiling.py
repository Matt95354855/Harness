import hashlib
import zipfile

import pytest
import pymupdf
from docx import Document

from mass_classification.profiling import profile_document
from mass_classification.extract import ExtractionError


def test_pdf_structure_and_hash_without_text_disclosure(tmp_path):
    path = tmp_path / 'case.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 50), 'CONFIDENTIAL native text')
        pdf.save(path)
    result = profile_document(path, path.name)
    assert result['pages'] == 1 and result['native_text_present']
    assert result['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result['tables'] is None and result['ocr_quality'] == 'not_measured'
    assert 'CONFIDENTIAL' not in str(result)


def test_docx_table_count_and_unknown_pagination(tmp_path):
    path = tmp_path / 'case.docx'
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = '1250 EUR'
    document.save(path)
    result = profile_document(path, path.name)
    assert result['tables'] == 1 and result['native_text_present']
    assert result['pages'] is None


def test_encrypted_pdf_and_invalid_files_are_rejected(tmp_path):
    path = tmp_path / 'locked.pdf'
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
    with pytest.raises(ExtractionError, match='Encrypted'):
        profile_document(path, path.name)
    with pytest.raises(ExtractionError, match='size'):
        profile_document(path, path.name, max_bytes=1)
    path.write_bytes(b'%PDF corrupt')
    with pytest.raises(ExtractionError):
        profile_document(path, path.name)


def test_docx_rejects_entities_and_fake_containers(tmp_path):
    path = tmp_path / 'bad.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY a "boom">]><x/>')
    with pytest.raises(ExtractionError, match='Unsafe'):
        profile_document(path, path.name)
    path.write_bytes(b'not a zip')
    with pytest.raises(ExtractionError):
        profile_document(path, path.name)
