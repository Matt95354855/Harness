"""Create non-sensitive mixed-content documents for the integration test."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Inches

root = Path('fixtures')
root.mkdir(exist_ok=True)
image = Image.new('RGB', (1400, 240), 'white')
font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 64)
ImageDraw.Draw(image).text((30, 70), 'ORION IMAGE FACTURE', font=font, fill='black')
image.save(root / 'image.png')
pdf = canvas.Canvas(str(root / 'sample.pdf'))
pdf.drawString(40, 780, 'Synthetic banking transaction virement')
pdf.drawImage(str(root / 'image.png'), 40, 550, width=520, height=90)
pdf.drawString(40, 480, 'DATE | MONTANT')
pdf.drawString(40, 450, '2026-04-14 | 1250 EUR')
pdf.save()
doc = Document()
doc.add_paragraph('Synthetic banking transaction virement')
doc.add_picture(str(root / 'image.png'), width=Inches(6))
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'DATE'
table.cell(0, 1).text = 'MONTANT'
table.cell(1, 0).text = '2026-04-14'
table.cell(1, 1).text = '1250 EUR'
doc.save(root / 'sample.docx')
